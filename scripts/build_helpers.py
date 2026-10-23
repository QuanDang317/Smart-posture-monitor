"""
Script chính: build_full_report.py
Tự động biên dịch toàn bộ nội dung học thuật, hình ảnh, bảng biểu và phụ lục kỹ thuật
thành 2 ấn bản báo cáo chuyên nghiệp:
1. File Word DOCX: docs/BAO_CAO_BAI_TAP_LON_SMART_POSTURE_MONITOR.docx
   (Đồng thời copy ra Desktop: c:/Users/win10/Desktop/BAO_CAO_BAI_TAP_LON_SMART_POSTURE_MONITOR.docx)
2. File Markdown: docs/BAO_CAO_BAI_TAP_LON_SMART_POSTURE_MONITOR.md
"""

import os
import sys
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

from report_data import *
from report_appendices import *

sys.stdout.reconfigure(encoding='utf-8')

# Colors palette
COLOR_PRIMARY_NAVY = "1B365D"
COLOR_SECONDARY_SLATE = "2C3E50"
COLOR_ACCENT_BLUE = "0066CC"
COLOR_BORDER_GREY = "D0D7DE"
COLOR_BG_LIGHT_GREY = "F8F9FA"
COLOR_BG_HEADER = "1B365D"
COLOR_CALLOUT_BG = "F0F7FF"

RGB_PRIMARY_NAVY = RGBColor(27, 54, 93)
RGB_SECONDARY_SLATE = RGBColor(44, 62, 80)
RGB_DARK_GREY = RGBColor(52, 73, 94)
RGB_TEXT_BLACK = RGBColor(20, 20, 20)

# Helper functions for Word Docx
def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'  <w:top w:w="{top}" w:type="dxa"/>'
        f'  <w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'  <w:left w:w="{left}" w:type="dxa"/>'
        f'  <w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="D0D7DE", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_styled_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.bold = True
    
    if level == 1:
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.2
        run.font.size = Pt(15)
        run.font.color.rgb = RGB_PRIMARY_NAVY
    elif level == 2:
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.2
        run.font.size = Pt(13.5)
        run.font.color.rgb = RGB_SECONDARY_SLATE
    elif level == 3:
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.2
        run.font.size = Pt(12.5)
        run.font.italic = True
        run.font.color.rgb = RGB_DARK_GREY
    return p

def add_body_paragraph(doc, text, bold_prefix="", italic=False, space_after=4.0, first_indent=0.5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.3
    if first_indent > 0:
        p.paragraph_format.first_line_indent = Inches(first_indent)
        
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(12.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGB_TEXT_BLACK
        
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12.5)
    run.font.italic = italic
    run.font.color.rgb = RGB_TEXT_BLACK
    return p

def add_bullet_point(doc, text, bold_title=""):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.25
    if bold_title:
        r_pre = p.add_run(bold_title)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(12.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGB_TEXT_BLACK
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12.5)
    run.font.color.rgb = RGB_TEXT_BLACK
    return p

def add_caption(doc, text, is_table=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if is_table:
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
    else:
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = is_table
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.italic = True
    run.font.color.rgb = RGB_SECONDARY_SLATE
    return p

def add_callout_box(doc, title, text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, COLOR_CALLOUT_BG)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Custom left thick border
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="{COLOR_PRIMARY_NAVY}"/>'
        f'  <w:top w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.2
    
    if title:
        r_t = p.add_run(title + "\n")
        r_t.font.name = "Times New Roman"
        r_t.font.size = Pt(11.5)
        r_t.font.bold = True
        r_t.font.color.rgb = RGB_PRIMARY_NAVY
        
    r_body = p.add_run(text)
    r_body.font.name = "Times New Roman"
    r_body.font.size = Pt(11)
    r_body.font.color.rgb = RGB_TEXT_BLACK
    
    # Empty space after table
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(4)

def add_code_block(doc, code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F4F6F8")
    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:left w:val="single" w:sz="6" w:space="0" w:color="CCCCCC"/>'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="CCCCCC"/>'
        f'  <w:right w:val="single" w:sz="6" w:space="0" w:color="CCCCCC"/>'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="CCCCCC"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(code_text.strip())
    run.font.name = "Consolas"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(30, 30, 30)
    
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_after = Pt(4)

def create_table_from_data(doc, table_data, col_widths=None, align_center_cols=None):
    rows_cnt = len(table_data)
    cols_cnt = len(table_data[0])
    
    tbl = doc.add_table(rows=rows_cnt, cols=cols_cnt)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl, color="D0D7DE", sz="4")
    
    # Header row
    for col_idx in range(cols_cnt):
        cell = tbl.cell(0, col_idx)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_background(cell, COLOR_BG_HEADER)
        set_cell_margins(cell, top=140, bottom=140, left=160, right=160)
        
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(table_data[0][col_idx])
        run.font.name = "Times New Roman"
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        
    # Repeat header on new page
    trPr = tbl.rows[0]._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

    # Data rows
    for row_idx in range(1, rows_cnt):
        bg_col = COLOR_BG_LIGHT_GREY if row_idx % 2 == 1 else "FFFFFF"
        for col_idx in range(cols_cnt):
            cell = tbl.cell(row_idx, col_idx)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_background(cell, bg_col)
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            
            p = cell.paragraphs[0]
            # Alignment
            if align_center_cols and col_idx in align_center_cols:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(table_data[row_idx][col_idx])
            run.font.name = "Times New Roman"
            run.font.size = Pt(10.5)
            run.font.color.rgb = RGB_TEXT_BLACK
            
    # Apply column widths if provided
    if col_widths and len(col_widths) == cols_cnt:
        for row in tbl.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)
                
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_after = Pt(4)

def add_image_safely(doc, image_path, caption_text, width_inches=5.8):
    if os.path.exists(image_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run()
        run.add_picture(image_path, width=Inches(width_inches))
        add_caption(doc, caption_text, is_table=False)
    else:
        print(f"Warning: Image not found: {image_path}")

print("Helper definitions complete.")
